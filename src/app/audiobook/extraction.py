"""Versioned source extraction. TTS normalization never runs in this module."""
from __future__ import annotations

import io
import posixpath
import re
import zipfile
from html.parser import HTMLParser
from pathlib import PurePosixPath
from urllib.parse import unquote
from typing import Any
from xml.etree import ElementTree as ET

from .hashing import bytes_hash, object_hash, text_hash
from .integrity import validate_revision
from .models import CanonicalChapter, SourceRevision
from .spans import UnicodeDialogueDetector


EXTRACTOR_VERSION = "audiobook-extractor-v9"
MAX_SOURCE_BYTES = 200 * 1024 * 1024
MAX_EPUB_UNCOMPRESSED_BYTES = 500 * 1024 * 1024
SUPPORTED_SOURCE_FORMATS = frozenset({
    "docx", "epub", "html", "htm", "markdown", "md", "pdf", "text", "txt",
})
_CHAPTER_HEADING = re.compile(r"^(?:#{1,2}\s+.+|chapter\s+(?:\d+|[IVXLCDM]+)\b.*)$", re.IGNORECASE)
_CHAPTER_WORD_HEADING = re.compile(r"^chapter\s+([a-z]+)\b.*$", re.IGNORECASE)
_CHAPTER_NUMBER_WORDS = frozenset({
    "zero", "one", "two", "three", "four", "five", "six", "seven", "eight", "nine",
    "ten", "eleven", "twelve", "thirteen", "fourteen", "fifteen", "sixteen",
    "seventeen", "eighteen", "nineteen", "twenty", "thirty", "forty", "fifty",
    "sixty", "seventy", "eighty", "ninety", "hundred", "thousand",
    "first", "second", "third", "fourth", "fifth", "sixth", "seventh", "eighth",
    "ninth", "tenth", "eleventh", "twelfth", "thirteenth", "fourteenth", "fifteenth",
    "sixteenth", "seventeenth", "eighteenth", "nineteenth", "twentieth", "thirtieth",
    "fortieth", "fiftieth", "sixtieth", "seventieth", "eightieth", "ninetieth",
})
_PAGE_RANGE = re.compile(r"^(\d+)(?:\s*-\s*(\d+))?$")
_TOC_CHAPTER_ENTRY = re.compile(
    r"^chapter\s+(?:\d+|[IVXLCDM]+|[a-z]+)\b.*"
    r"(?:\.{2,}|\s{2,})\s*\d{1,4}\s*$",
    re.IGNORECASE,
)


class UnsupportedSource(ValueError):
    pass


def parse_page_ranges(value: str) -> list[list[int]]:
    """Parse a user-facing 1-based page range list into merged inclusive ranges."""
    ranges: list[tuple[int, int]] = []
    for item in value.split(","):
        token = item.strip()
        if not token:
            raise UnsupportedSource("excluded PDF pages must be comma-separated numbers or ranges")
        match = _PAGE_RANGE.fullmatch(token)
        if match is None:
            raise UnsupportedSource(
                "excluded PDF pages must use 1-based values such as 1-3, 42-45"
            )
        start = int(match.group(1))
        end = int(match.group(2) or start)
        if start < 1 or end < start:
            raise UnsupportedSource("excluded PDF page ranges must start at 1 and end after their start")
        ranges.append((start, end))

    merged: list[list[int]] = []
    for start, end in sorted(ranges):
        if merged and start <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return merged


def normalize_extraction_settings(
    source_format: str, settings: dict[str, object] | None = None,
) -> dict[str, object]:
    """Validate and canonicalize settings before they become revision identity."""
    raw = dict(settings or {})
    if not raw:
        return {}
    if source_format != "pdf":
        raise UnsupportedSource("page exclusions are only supported for PDF sources")
    if set(raw) != {"excluded_page_ranges"}:
        raise UnsupportedSource("unknown extraction settings")
    raw_ranges = raw["excluded_page_ranges"]
    if not isinstance(raw_ranges, list):
        raise UnsupportedSource("excluded PDF pages must be a list of ranges")

    ranges: list[tuple[int, int]] = []
    for raw_range in raw_ranges:
        if not isinstance(raw_range, (list, tuple)) or len(raw_range) != 2:
            raise UnsupportedSource("excluded PDF pages must be inclusive two-number ranges")
        start, end = raw_range
        if (isinstance(start, bool) or not isinstance(start, int) or
                isinstance(end, bool) or not isinstance(end, int)):
            raise UnsupportedSource("excluded PDF pages must be inclusive two-number ranges")
        if start < 1 or end < start:
            raise UnsupportedSource("excluded PDF page ranges must start at 1 and end after their start")
        ranges.append((start, end))

    if not ranges:
        return {}
    merged: list[list[int]] = []
    for start, end in sorted(ranges):
        if merged and start <= merged[-1][1] + 1:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return {"excluded_page_ranges": merged}


class _ReadingHTML(HTMLParser):
    _BLOCKS = {"p", "div", "section", "article", "h1", "h2", "h3", "h4", "blockquote", "li"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.suppressed = 0
        self.headings: list[str] = []
        self.heading_positions: list[dict[str, object]] = []
        self._heading: list[str] | None = None
        self._heading_start: int | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "body":
            self.suppressed = 0
        if tag in {"script", "style", "nav", "head"}:
            self.suppressed += 1
            return
        if self.suppressed:
            return
        if tag in self._BLOCKS and self.parts and not self.parts[-1].endswith("\n"):
            self.parts.append("\n\n")
        if tag == "br":
            self.parts.append("\n")
        if tag in {"h1", "h2", "h3"}:
            self._heading = []
            self._heading_start = sum(len(part) for part in self.parts)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"script", "style", "nav", "head"}:
            self.suppressed = max(0, self.suppressed - 1)
            return
        if self.suppressed:
            return
        if tag in {"h1", "h2", "h3"} and self._heading is not None:
            title = "".join(self._heading).strip()
            if title:
                self.headings.append(title)
                self.heading_positions.append({
                    "text": title,
                    "start_offset": int(self._heading_start or 0),
                })
            self._heading = None
            self._heading_start = None
        if tag in self._BLOCKS:
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if not self.suppressed:
            self.parts.append(data)
            if self._heading is not None:
                self._heading.append(data)


def _decode_utf8(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
    except UnicodeDecodeError as exc:
        raise UnsupportedSource("source must be valid UTF-8") from exc


def _is_chapter_heading(line: str) -> bool:
    # TOC entries often begin with "Chapter ..." but are not real narrative
    # boundaries. Keep them inside the surrounding source region so the
    # document-structure stage can classify and skip them coherently.
    if _TOC_CHAPTER_ENTRY.match(line):
        return False
    if _CHAPTER_HEADING.match(line):
        return True
    match = _CHAPTER_WORD_HEADING.match(line)
    return bool(match and match.group(1).casefold() in _CHAPTER_NUMBER_WORDS)


def _text_chapters(content: str) -> list[tuple[str, str]]:
    chapters: list[tuple[str, str]] = []
    title = "Opening"
    current: list[str] = []
    for line in content.splitlines(keepends=True):
        stripped = line.strip()
        is_heading = _is_chapter_heading(stripped)
        if is_heading and current:
            chapters.append((title, "".join(current)))
            current = []
        if is_heading:
            title = stripped.lstrip("# ") or title
        current.append(line)
    if current:
        chapters.append((title, "".join(current)))
    return chapters


def _epub_chapters(content: bytes) -> tuple[list[tuple[str, str]], dict[str, Any], list[str]]:
    chapters: list[tuple[str, str]] = []
    warnings: list[str] = []
    metadata: dict[str, Any] = {}
    semantic_headings: list[dict[str, object]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_EPUB_UNCOMPRESSED_BYTES:
                raise UnsupportedSource("EPUB uncompressed content exceeds the supported limit")
            names = set(archive.namelist())
            if "META-INF/encryption.xml" in names:
                raise UnsupportedSource("encrypted EPUB is unsupported")
            container = ET.fromstring(archive.read("META-INF/container.xml"))
            rootfile = container.find(".//{*}rootfile")
            if rootfile is None or not rootfile.get("full-path"):
                raise UnsupportedSource("EPUB container has no package document")
            opf_path = rootfile.attrib["full-path"]
            opf = ET.fromstring(archive.read(opf_path))
            for key in ("title", "creator", "language"):
                node = opf.find(f".//{{*}}metadata/{{*}}{key}")
                if node is not None and node.text:
                    metadata[key] = node.text.strip()
            manifest: dict[str, tuple[str, str, str]] = {}
            for item in opf.findall(".//{*}manifest/{*}item"):
                item_id = item.get("id")
                if item_id:
                    manifest[item_id] = (item.get("href", ""), item.get("media-type", ""), item.get("properties", ""))
            base = str(PurePosixPath(opf_path).parent)
            for spine_item in opf.findall(".//{*}spine/{*}itemref"):
                item = manifest.get(spine_item.get("idref", ""))
                if not item:
                    raise UnsupportedSource("EPUB spine item missing from manifest")
                href, media_type, properties = item
                if "nav" in properties.split() or media_type not in {"application/xhtml+xml", "text/html"}:
                    continue
                path = posixpath.normpath(posixpath.join(base, unquote(href.split("#", 1)[0])))
                if path not in names:
                    raise UnsupportedSource(f"EPUB spine resource missing: {path}")
                parser = _ReadingHTML()
                parser.feed(_decode_utf8(archive.read(path)))
                raw_text = "".join(parser.parts)
                leading_trim = len(raw_text) - len(raw_text.lstrip("\n"))
                text = raw_text.strip("\n")
                if text.strip():
                    chapter_index = len(chapters)
                    chapters.append((parser.headings[0] if parser.headings else f"Chapter {chapter_index + 1}", text))
                    semantic_headings.extend(
                        {
                            "chapter_index": chapter_index,
                            "heading_index": heading_index,
                            "text": str(heading.get("text") or ""),
                            "start_offset": max(
                                0, int(heading.get("start_offset") or 0) - leading_trim
                            ),
                        }
                        for heading_index, heading in enumerate(parser.heading_positions)
                    )
            if not chapters:
                raise UnsupportedSource("EPUB has no readable spine chapters")
            metadata["epub_semantic_headings"] = semantic_headings
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        raise UnsupportedSource("invalid EPUB package") from exc
    return chapters, metadata, warnings


def _html_chapters(content: bytes) -> tuple[list[tuple[str, str]], dict[str, Any], list[str]]:
    parser = _ReadingHTML()
    try:
        parser.feed(_decode_utf8(content))
        parser.close()
    except Exception as exc:
        raise UnsupportedSource("invalid HTML document") from exc
    raw_text = "".join(parser.parts)
    leading_trim = len(raw_text) - len(raw_text.lstrip())
    text = raw_text.strip()
    if not text:
        raise UnsupportedSource("HTML document has no readable text")
    chapters = _text_chapters(text)
    chapter_ranges: list[tuple[int, int]] = []
    cursor = 0
    for _title, chapter_text in chapters:
        chapter_ranges.append((cursor, cursor + len(chapter_text)))
        cursor += len(chapter_text)
    semantic_headings: list[dict[str, object]] = []
    for heading_index, heading in enumerate(parser.heading_positions):
        absolute = max(0, int(heading.get("start_offset") or 0) - leading_trim)
        for chapter_index, (start, end) in enumerate(chapter_ranges):
            if start <= absolute < end:
                semantic_headings.append({
                    "chapter_index": chapter_index,
                    "heading_index": heading_index,
                    "text": str(heading.get("text") or ""),
                    "start_offset": absolute - start,
                })
                break
    return chapters, {
        "html_semantic_headings": semantic_headings,
    }, []


def _pdf_outline_chapters(reader: object, pages: list[tuple[int, str]]) -> list[tuple[str, str]]:
    """Use page-aligned PDF bookmarks when they identify real headings."""
    included = {number: text for number, text in pages}
    markers: dict[int, str] = {}

    def walk(items: object) -> None:
        if not isinstance(items, (list, tuple)):
            return
        for item in items:
            if isinstance(item, (list, tuple)):
                walk(item)
                continue
            try:
                title = str(item.title).strip()
                number = reader.get_destination_page_number(item) + 1
            except (AttributeError, KeyError, ValueError):
                continue
            text = included.get(number)
            if not title or text is None:
                continue
            heading = re.sub(r"\W+", "", title).casefold()
            opening = re.sub(r"\W+", "", " ".join(text.splitlines()[:3])).casefold()
            if heading and heading in opening:
                markers.setdefault(number, title)

    try:
        walk(reader.outline)
    except Exception:
        return []
    if len(markers) < 2:
        return []
    chapters: list[tuple[str, str]] = []
    title = "Opening"
    current: list[str] = []
    for number, text in pages:
        if number in markers:
            if current:
                chapters.append((title, "\n\n".join(current)))
            title, current = markers[number], []
        current.append(text)
    if current:
        chapters.append((title, "\n\n".join(current)))
    return chapters


def _pdf_chapters(
    content: bytes, *, settings: dict[str, object],
) -> tuple[list[tuple[str, str]], dict[str, Any], list[str]]:
    try:
        from PyPDF2 import PdfReader
    except ImportError as exc:  # pragma: no cover - packaging failure
        raise UnsupportedSource("PDF support requires PyPDF2") from exc
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
    except Exception as exc:
        raise UnsupportedSource("invalid PDF document") from exc

    page_count = len(reader.pages)
    excluded_ranges = settings.get("excluded_page_ranges", [])
    for raw_range in excluded_ranges:
        start, end = raw_range
        if end > page_count:
            raise UnsupportedSource(
                f"excluded PDF page range {start}-{end} exceeds this document's {page_count} pages"
            )

    pages: list[tuple[int, str]] = []
    page_edges: list[dict[str, object]] = []
    page_blocks: list[dict[str, object]] = []
    for page_number, page in enumerate(reader.pages, start=1):
        if any(start <= page_number <= end for start, end in excluded_ranges):
            continue
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise UnsupportedSource("PDF text extraction failed") from exc
        if text.strip():
            clean_text = text.strip()
            pages.append((page_number, clean_text))
            lines = [line.strip() for line in clean_text.splitlines() if line.strip()]
            page_edges.append({
                "page_index": page_number - 1,
                "top": lines[:3],
                "bottom": lines[-3:],
            })
            denominator = max(1, len(lines) - 1)
            page_blocks.extend(
                {
                    "page_index": page_number - 1,
                    "block_index": line_index,
                    "reading_order": line_index,
                    "original_text": line,
                    "normalized_text": " ".join(line.split()).casefold(),
                    "distance_from_top": line_index / denominator,
                    "distance_from_bottom": (len(lines) - 1 - line_index) / denominator,
                    "bounding_box": None,
                    "font_size": None,
                    "font_weight": None,
                    "font_style": None,
                }
                for line_index, line in enumerate(lines)
            )
    if not pages:
        if excluded_ranges:
            raise UnsupportedSource("page exclusions removed all readable PDF pages")
        raise UnsupportedSource("PDF has no extractable text; scanned PDFs are unsupported")

    metadata: dict[str, Any] = {}
    for key, value in (reader.metadata or {}).items():
        if value is not None and str(value).strip():
            metadata[str(key).lstrip("/").lower()] = str(value).strip()
    metadata["pdf_page_edges"] = page_edges
    metadata["pdf_page_blocks"] = page_blocks
    metadata["pdf_page_count"] = page_count
    warnings = []
    if excluded_ranges:
        rendered_ranges = ", ".join(
            str(start) if start == end else f"{start}-{end}"
            for start, end in excluded_ranges
        )
        warnings.append(f"Excluded PDF pages: {rendered_ranges}")
    chapters = _pdf_outline_chapters(reader, pages)
    if not chapters:
        chapters = _text_chapters("\n\n".join(text for _, text in pages))
        # A title page may share the first physical page with Chapter One. Keep
        # that text, but do not expose the short title-page fragment as a
        # separate chapter when the first page already contains a real heading.
        if len(chapters) > 1 and chapters[0][0] == "Opening":
            first_page_has_heading = any(
                _is_chapter_heading(line.strip()) for line in pages[0][1].splitlines()
            )
            if first_page_has_heading:
                first_title, first_text = chapters[1]
                chapters = [(first_title, chapters[0][1] + first_text), *chapters[2:]]
    return chapters, metadata, warnings


def _docx_chapters(content: bytes) -> tuple[list[tuple[str, str]], dict[str, Any], list[str]]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - packaging failure
        raise UnsupportedSource("DOCX support requires python-docx") from exc
    try:
        document = Document(io.BytesIO(content))
    except Exception as exc:
        raise UnsupportedSource("invalid DOCX document") from exc

    blocks: list[str] = []
    style_blocks: list[dict[str, object]] = []
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        block_index = len(blocks)
        blocks.append(paragraph.text)
        runs = [run for run in paragraph.runs if run.text.strip()]
        style_blocks.append({
            "block_index": block_index,
            "text": paragraph.text,
            "normalized_text": " ".join(paragraph.text.split()).casefold(),
            "style": str(getattr(paragraph.style, "name", "") or ""),
            "bold": bool(runs) and all(bool(run.bold) for run in runs),
            "italic": bool(runs) and all(bool(run.italic) for run in runs),
        })
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                text_row = " | ".join(cells)
                style_blocks.append({
                    "block_index": len(blocks),
                    "text": text_row,
                    "normalized_text": " ".join(text_row.split()).casefold(),
                    "style": "Table",
                    "bold": False,
                    "italic": False,
                })
                blocks.append(text_row)
    text = "\n\n".join(blocks).strip()
    if not text:
        raise UnsupportedSource("DOCX document has no readable text")

    metadata: dict[str, Any] = {"docx_style_blocks": style_blocks}
    properties = document.core_properties
    if properties.title:
        metadata["title"] = properties.title.strip()
    if properties.author:
        metadata["creator"] = properties.author.strip()
    return _text_chapters(text), metadata, []


def extract_source(
    *, project_id: str, content: bytes, source_format: str,
    settings: dict[str, object] | None = None,
) -> SourceRevision:
    if source_format not in SUPPORTED_SOURCE_FORMATS:
        raise UnsupportedSource(f"unsupported source format: {source_format}")
    if len(content) > MAX_SOURCE_BYTES:
        raise UnsupportedSource("source exceeds the supported size limit")
    settings = normalize_extraction_settings(source_format, settings)
    if source_format == "epub":
        source_chapters, metadata, warnings = _epub_chapters(content)
    elif source_format == "pdf":
        source_chapters, metadata, warnings = _pdf_chapters(content, settings=settings)
    elif source_format == "docx":
        source_chapters, metadata, warnings = _docx_chapters(content)
    elif source_format in {"html", "htm"}:
        source_chapters, metadata, warnings = _html_chapters(content)
    else:
        source_chapters = _text_chapters(_decode_utf8(content))
        metadata, warnings = {}, []
    if not source_chapters or not any(text.strip() for _, text in source_chapters):
        raise UnsupportedSource("source has no readable text")
    original_hash = bytes_hash(content)
    settings_hash = object_hash(settings)
    chapter_hashes = [text_hash(text) for _, text in source_chapters]
    detector = UnicodeDialogueDetector()
    canonical_hash = object_hash({
        "extractor_version": EXTRACTOR_VERSION,
        "settings_hash": settings_hash,
        "original_asset_hash": original_hash,
        "chapters": chapter_hashes,
    })
    # Span segmentation is part of revision identity without changing the
    # canonical-text hash contract. A detector upgrade therefore creates a new
    # durable revision while integrity validation remains about source content.
    revision_id = (
        f"ab:sr:{text_hash(f'{project_id}:{canonical_hash}:{detector.version}')}"
    )
    chapters = tuple(
        CanonicalChapter(
            id=f"ab:ch:{text_hash(f'{revision_id}:{ordinal}:{chapter_hashes[ordinal]}')}",
            ordinal=ordinal,
            title=title,
            canonical_text=text,
            canonical_hash=chapter_hashes[ordinal],
            spans=(),
            structure={"source_format": source_format},
        )
        for ordinal, (title, text) in enumerate(source_chapters)
    )
    chapters = tuple(
        CanonicalChapter(
            id=chapter.id, ordinal=chapter.ordinal, title=chapter.title,
            canonical_text=chapter.canonical_text, canonical_hash=chapter.canonical_hash,
            spans=detector.detect(chapter.id, chapter.canonical_text), structure=chapter.structure,
        )
        for chapter in chapters
    )
    revision = SourceRevision(
        id=revision_id, project_id=project_id,
        original_asset_hash=original_hash, source_format=source_format,
        extractor_version=EXTRACTOR_VERSION, extraction_settings=settings,
        extraction_settings_hash=settings_hash, canonical_hash=canonical_hash,
        chapters=chapters, metadata=metadata, warnings=tuple(warnings),
    )
    validate_revision(revision)
    return revision


def resegment_revision(
    revision: SourceRevision, *, styles: tuple[str, ...],
    discovery: dict[str, object],
) -> SourceRevision:
    """Apply a verified style without changing any extracted source text."""
    detector = UnicodeDialogueDetector(styles=styles)
    style_identity = ",".join(detector.styles)
    revision_id = f"ab:sr:{text_hash(f'{revision.project_id}:{revision.canonical_hash}:{detector.version}:{style_identity}')}"
    chapters = tuple(
        CanonicalChapter(
            id=f"ab:ch:{text_hash(f'{revision_id}:{chapter.ordinal}:{chapter.canonical_hash}')}",
            ordinal=chapter.ordinal,
            title=chapter.title,
            canonical_text=chapter.canonical_text,
            canonical_hash=chapter.canonical_hash,
            spans=(),
            structure=chapter.structure,
        )
        for chapter in revision.chapters
    )
    chapters = tuple(
        CanonicalChapter(
            id=chapter.id, ordinal=chapter.ordinal, title=chapter.title,
            canonical_text=chapter.canonical_text, canonical_hash=chapter.canonical_hash,
            spans=detector.detect(chapter.id, chapter.canonical_text),
            structure=chapter.structure,
        )
        for chapter in chapters
    )
    updated = SourceRevision(
        id=revision_id, project_id=revision.project_id,
        original_asset_hash=revision.original_asset_hash,
        source_format=revision.source_format,
        extractor_version=revision.extractor_version,
        extraction_settings=revision.extraction_settings,
        extraction_settings_hash=revision.extraction_settings_hash,
        canonical_hash=revision.canonical_hash,
        chapters=chapters,
        metadata={**revision.metadata, "dialogue_style_discovery": discovery},
        warnings=revision.warnings,
    )
    validate_revision(updated)
    return updated
